"""Rotas e regras da área: relatorios."""
from ..templates_loader import ler_template
import io
from datetime import date, timedelta
from xml.sax.saxutils import escape as xml_escape
from flask import abort, current_app, flash, redirect, request, send_file
from ..base import agora, auditar, db, hoje, permissao, render_page, tem_permissao
from ..principal import bp

@bp.route('/relatorios')
@permissao('relatorios')
def relatorios():
    available = {key: report for key, report in current_app.config['SAGEL_REPORTS'].items() if tem_permissao(report['modulo'], 'gerenciar')}
    key = request.args.get('tipo', '')
    report = available.get(key)
    rows, columns = ([], [])
    if key and (not report):
        abort(403)
    start, end = (request.args.get('inicio', ''), request.args.get('fim', ''))
    try:
        if start:
            date.fromisoformat(start)
        if end:
            date.fromisoformat(end)
        if start and end and (start > end):
            raise ValueError()
    except ValueError:
        flash('Informe um período válido.', 'erro')
        return redirect('/relatorios')
    if report:
        query = 'SELECT * FROM (' + report['sql'].strip().rstrip(';') + ') rel'
        args, where = ([], [])
        column = report.get('data_coluna')
        if column:
            quoted = '"' + column.replace('"', '""') + '"'
            if start:
                where.append(quoted + '>=?')
                args.append(start)
            if end:
                where.append(quoted + '<?')
                args.append((date.fromisoformat(end) + timedelta(days=1)).isoformat())
        if where:
            query += ' WHERE ' + ' AND '.join(where)
        cursor = db().execute(query + ' LIMIT 10000', args)
        columns = [c[0] for c in cursor.description]
        rows = cursor.fetchall()
        formato = request.args.get('formato')
        if formato in ('xlsx', 'pdf'):
            auditar('relatorios', 'Exportação ' + formato, key, f'{start} a {end}; {len(rows)} registros')
            db().commit()
            if formato == 'xlsx':
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill
                from openpyxl.utils import get_column_letter
                book = Workbook()
                sheet = book.active
                sheet.title = 'Relatório'
                sheet.append(columns)
                for row in rows:
                    sheet.append(list(row))
                    for cell in sheet[sheet.max_row]:
                        if isinstance(cell.value, str):
                            cell.data_type = 's'
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='00663B')
                for i, column_name in enumerate(columns, 1):
                    sheet.column_dimensions[get_column_letter(i)].width = min(45, max(16, len(column_name) + 3))
                sheet.freeze_panes = 'A2'
                sheet.auto_filter.ref = sheet.dimensions
                output = io.BytesIO()
                book.save(output)
                mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            else:
                from reportlab.lib import colors
                from reportlab.lib.pagesizes import A4, A3, landscape
                from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
                output = io.BytesIO()
                size = landscape(A3 if len(columns) > 8 else A4)
                document = SimpleDocTemplate(output, pagesize=size, rightMargin=28, leftMargin=28, topMargin=30, bottomMargin=30)
                style = getSampleStyleSheet()
                compact = ParagraphStyle('Cell', parent=style['BodyText'], fontSize=8, leading=10, wordWrap='CJK')
                data = [[Paragraph(xml_escape(str(c)), compact) for c in columns]]
                for row in rows:
                    data.append([Paragraph(xml_escape(str(value if value is not None else '')[:350]), compact) for value in row])
                table = LongTable(data, repeatRows=1, colWidths=[(size[0] - 56) / len(columns)] * len(columns))
                table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eaf4ee')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#dce6df')), ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))

                def rodape(canvas, doc):
                    canvas.setFont('Helvetica', 8)
                    canvas.drawString(28, 16, 'SAGEL | ' + agora()[:10])
                    canvas.drawRightString(size[0] - 28, 16, f'Página {doc.page}')
                document.build([Paragraph('SAGEL — ' + xml_escape(report['titulo']), style['Title']), Paragraph(xml_escape(f"Período: {start or 'início'} a {end or 'hoje'} | {len(rows)} registros"), style['Normal']), Spacer(1, 14), table], onFirstPage=rodape, onLaterPages=rodape)
                mime = 'application/pdf'
            output.seek(0)
            return send_file(output, as_attachment=True, download_name=f'sagel-{key}-{hoje()}.{formato}', mimetype=mime)
    return render_page('Relatórios e indicadores', ler_template('relatorios/relatorios.html'), available=available, key=key, report=report, rows=rows, columns=columns)
