"""Regressões de consultas portáveis, sempre com banco e arquivos temporários."""
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
import sqlite3
import shutil
import unittest
import uuid
from unittest.mock import patch

from openpyxl import load_workbook

from sagel import create_app
from sagel.base import db
from sagel.frota import indicadores


class PortabilidadeRotasTests(unittest.TestCase):
    def setUp(self):
        self.folder = Path(__file__).resolve().parent / ('sagel-portabilidade-' + uuid.uuid4().hex)
        self.folder.mkdir()
        self.addCleanup(self.cleanup_folder)
        self.database = self.folder / 'teste.db'
        self.app = create_app({'TESTING': True, 'DATABASE': str(self.database),
                               'SECRET_KEY': 'teste-portabilidade-isolado',
                               'STORAGE_DIR': str(self.folder / 'anexos'),
                               'BACKUP_DIR': str(self.folder / 'backups')})
        self.client = self.app.test_client()
        with self.app.app_context():
            db().execute("""INSERT INTO usuarios(id,nome,usuario,email,senha,cargo,admin,perfil)
                          VALUES(1,'Gestor','gestor','gestor@teste.local','invalida','Gestor',1,'admin')""")
            db().execute("""INSERT INTO frota_veiculos
                          (id,placa,modelo,ano,tipo_combustivel,quilometragem,criado_em)
                          VALUES(1,'ABC1D23','Modelo Especial',2024,'Diesel S10',100,'2030-01-01')""")
            db().commit()
        with self.client.session_transaction() as session:
            session['usuario_id'] = 1
            session['csrf_token'] = 'token-teste'

    def cleanup_folder(self):
        if self.folder.parent == Path(__file__).resolve().parent and self.folder.name.startswith('sagel-portabilidade-'):
            shutil.rmtree(self.folder)

    def post(self, url, data):
        return self.client.post(url, data={'csrf_token': 'token-teste', **data})

    def test_documentos_incluem_trigesimo_dia_sem_incluir_seguinte(self):
        today = '2030-02-01'
        cutoff = date.fromisoformat(today) + timedelta(days=30)
        with self.app.app_context():
            for title, expiry in (('Vencido teste', today), ('Limite teste', cutoff.isoformat()),
                                  ('Futuro teste', (cutoff + timedelta(days=1)).isoformat())):
                db().execute("""INSERT INTO frota_documentos
                              (veiculo_id,titulo,validade,chave,nome_arquivo,autor_id,criado_em)
                              VALUES(1,?,?,?,'teste.pdf',1,?)""", (title, expiry, title, today))
            db().commit()
            with patch('sagel.frota.hoje', return_value=today):
                self.assertEqual(indicadores(db())[-1]['valor'], 2)
        with patch('sagel.modulos.veiculos.hoje', return_value=today):
            response = self.client.get('/frota')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Vencido teste', response.text)
        self.assertIn('Limite teste', response.text)
        self.assertNotIn('Futuro teste', response.text)

    def test_busca_ignora_caixa_e_placa_duplicada_e_rejeitada(self):
        response = self.client.get('/frota?q=mODELo eSPECIAL')
        self.assertEqual(response.status_code, 200)
        self.assertIn('ABC1D23', response.text)
        with self.app.app_context():
            with self.assertRaises(sqlite3.IntegrityError):
                db().execute("""INSERT INTO frota_veiculos
                              (placa,modelo,ano,tipo_combustivel,criado_em)
                              VALUES('abc1d23','Duplicado',2024,'Diesel S10','2030-01-01')""")
            db().rollback()
            self.assertEqual(db().execute('SELECT count(*) FROM frota_veiculos').fetchone()[0], 1)

    def test_configuracoes_atualizadas_sem_duplicar_registro(self):
        self.assertEqual(self.post('/configuracoes', {'cadastro_aberto': '1'}).status_code, 302)
        self.assertEqual(self.post('/configuracoes', {'backup_automatico': '1'}).status_code, 302)
        with self.app.app_context():
            settings = {row['chave']: row['valor'] for row in db().execute('SELECT * FROM configuracoes')}
        self.assertEqual(settings['cadastro_aberto'], '0')
        self.assertEqual(settings['backup_automatico'], '1')

    def test_relatorio_filtra_periodo_e_exporta_valores(self):
        with self.app.app_context():
            for title, deadline in (('Tarefa do periodo', '2030-03-01'), ('Tarefa posterior', '2030-03-02')):
                db().execute("""INSERT INTO tarefas(titulo,responsavel,responsavel_id,criado_por,prazo,data)
                              VALUES(?,'Gestor',1,1,?,'2030-02-01')""", (title, deadline))
            db().commit()
        params = {'tipo': 'tarefas', 'inicio': '2030-03-01', 'fim': '2030-03-01'}
        response = self.client.get('/relatorios', query_string=params)
        self.assertEqual(response.status_code, 200)
        self.assertIn('Tarefa do periodo', response.text)
        self.assertNotIn('Tarefa posterior', response.text)
        response = self.client.get('/relatorios', query_string={**params, 'formato': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.data)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['A1'].value, 'Tarefa')
        self.assertEqual(sheet['A2'].value, 'Tarefa do periodo')
        self.assertEqual(sheet['D2'].value, '2030-03-01')
